"""Document creation service with Word/PDF output."""
import json
from pathlib import Path
from typing import Dict, List, Optional
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from ..config import settings

# Citation styles
CITATION_STYLES = {
    "Vancouver": {
        "format": "{authors}. {title}. {journal}. {year};{volume}:{pages}.",
        "in_text": "[{number}]"
    },
    "APA": {
        "format": "{authors} ({year}). {title}. {journal}, {volume}, {pages}.",
        "in_text": "({author}, {year})"
    },
    "Harvard": {
        "format": "{authors} ({year}) '{title}', {journal}, {volume}, pp. {pages}.",
        "in_text": "({author} {year})"
    }
}


def format_citation(paper: Dict, style: str = "Vancouver", number: int = 1) -> str:
    """Format a citation according to style.

    Args:
        paper: Paper metadata
        style: Citation style name
        number: Citation number

    Returns:
        Formatted citation string
    """
    authors = paper.get("authors", [])
    if len(authors) > 3:
        author_str = f"{authors[0]} et al."
    elif authors:
        author_str = ", ".join(authors)
    else:
        author_str = "Unknown"

    template = CITATION_STYLES.get(style, CITATION_STYLES["Vancouver"])["format"]

    return template.format(
        authors=author_str,
        title=paper.get("title", "Unknown"),
        journal=paper.get("journal", ""),
        year=paper.get("year", ""),
        volume=paper.get("volume", ""),
        pages=paper.get("pages", "")
    )


def create_word_document(
    content: Dict,
    output_path: str,
    citation_style: str = "Vancouver"
) -> str:
    """Create a Word document.

    Args:
        content: Document content with sections and citations
        output_path: Output file path
        citation_style: Citation format style

    Returns:
        Path to created file
    """
    doc = Document()

    # Title
    title = doc.add_heading(content.get("title", "Document"), 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Date
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para.add_run(datetime.now().strftime("%Y-%m-%d"))

    doc.add_paragraph()

    # Sections
    for section in content.get("sections", []):
        doc.add_heading(section.get("name", ""), level=1)
        para = doc.add_paragraph(section.get("content", ""))
        para.style.font.size = Pt(11)

    # References
    if content.get("citations"):
        doc.add_heading("References", level=1)
        for i, citation in enumerate(content["citations"], 1):
            ref_text = format_citation(citation, citation_style, i)
            para = doc.add_paragraph()
            para.add_run(f"[{i}] ").bold = True
            para.add_run(ref_text)

    # Save
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    return str(output_path)


def create_pdf_document(
    content: Dict,
    output_path: str,
    citation_style: str = "Vancouver"
) -> str:
    """Create a PDF document.

    Args:
        content: Document content with sections and citations
        output_path: Output file path
        citation_style: Citation format style

    Returns:
        Path to created file
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Title'],
        fontSize=18,
        spaceAfter=20,
        alignment=1  # Center
    )

    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading1'],
        fontSize=14,
        spaceAfter=10,
        spaceBefore=20
    )

    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['Normal'],
        fontSize=11,
        spaceAfter=10,
        leading=14
    )

    story = []

    # Title
    story.append(Paragraph(content.get("title", "Document"), title_style))
    story.append(Paragraph(datetime.now().strftime("%Y-%m-%d"), styles['Normal']))
    story.append(Spacer(1, 20))

    # Sections
    for section in content.get("sections", []):
        story.append(Paragraph(section.get("name", ""), heading_style))
        # Split content into paragraphs
        section_content = section.get("content", "")
        for para in section_content.split("\n\n"):
            if para.strip():
                story.append(Paragraph(para.strip(), body_style))

    # References
    if content.get("citations"):
        story.append(Paragraph("References", heading_style))
        for i, citation in enumerate(content["citations"], 1):
            ref_text = format_citation(citation, citation_style, i)
            story.append(Paragraph(f"[{i}] {ref_text}", body_style))

    doc.build(story)

    return str(output_path)


def create_document(
    content: Dict,
    output_path: str,
    file_type: str = "docx",
    citation_style: str = "Vancouver"
) -> str:
    """Create a document in specified format.

    Args:
        content: Document content
        output_path: Output file path
        file_type: 'docx' or 'pdf'
        citation_style: Citation format style

    Returns:
        Path to created file
    """
    if file_type == "pdf":
        return create_pdf_document(content, output_path, citation_style)
    else:
        return create_word_document(content, output_path, citation_style)
